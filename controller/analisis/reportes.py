"""Generacion de reportes JSON, texto y comparativos.

Autor: GPT-5.4
"""

from __future__ import annotations

import json
from datetime import datetime

import pandas as pd

from controller.utils.logger import get_logger

logger = get_logger(__name__)


def generar_reporte(resultado: dict) -> str:
    reporte = {
        "generado_en": datetime.now().isoformat(),
        "contrato_id": resultado.get("contrato_id"),
        "nombre": resultado.get("nombre"),
        "categoria": resultado.get("categoria"),
        "score": resultado.get("score"),
        "nivel_riesgo": resultado.get("nivel_riesgo"),
        "resumen_ejecutivo": resultado.get("resumen_ejecutivo"),
        "resumen": resultado.get("estadisticas"),
        "comparacion_dataset": resultado.get("comparacion_dataset"),
        "anomalias": resultado.get("anomalias"),
        "hallazgos": resultado.get("hallazgos"),
        "clausulas": resultado.get("clausulas"),
    }
    logger.info(f"Reporte generado para contrato {resultado.get('contrato_id')}")
    return json.dumps(reporte, ensure_ascii=False, indent=2)


def generar_reporte_texto(resultado: dict) -> str:
    hallazgos = resultado.get("hallazgos", [])
    bloques = [
        f"Reporte LexCore - {resultado.get('nombre')}",
        f"Contrato ID: {resultado.get('contrato_id')}",
        f"Categoria: {resultado.get('categoria')}",
        f"Score: {resultado.get('score')} / 100",
        f"Nivel de riesgo: {resultado.get('nivel_riesgo')}",
        "",
        resultado.get("resumen_ejecutivo", ""),
        "",
        "Hallazgos principales:",
    ]
    if hallazgos:
        for hallazgo in hallazgos:
            bloques.append(
                f"- Clausula {hallazgo.get('clausula_id')}: {hallazgo.get('tipo_riesgo')} "
                f"({hallazgo.get('severidad')}) | {hallazgo.get('articulo_violado')} | "
                f"{hallazgo.get('descripcion_riesgo')}"
            )
    else:
        bloques.append("- No se detectaron hallazgos de riesgo relevantes.")
    return "\n".join(bloques)

def generar_reporte_pdf(resultado: dict) -> bytes:
    from io import BytesIO

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(f"Reporte LexCore - {resultado.get('nombre', 'Contrato')}", estilos["Title"]),
        Spacer(1, 0.5 * cm),
        Paragraph(f"<b>Contrato ID:</b> {resultado.get('contrato_id')}", estilos["Normal"]),
        Paragraph(f"<b>Categoria:</b> {resultado.get('categoria')}", estilos["Normal"]),
        Paragraph(f"<b>Score:</b> {resultado.get('score')} / 100", estilos["Normal"]),
        Paragraph(f"<b>Nivel de riesgo:</b> {resultado.get('nivel_riesgo')}", estilos["Normal"]),
        Spacer(1, 0.5 * cm),
        Paragraph(resultado.get("resumen_ejecutivo", ""), estilos["Normal"]),
        Spacer(1, 0.7 * cm),
        Paragraph("Hallazgos principales", estilos["Heading2"]),
    ]

    hallazgos = resultado.get("hallazgos", [])
    if hallazgos:
        filas = [["Cláusula", "Riesgo", "Severidad", "Artículo"]]
        for hallazgo in hallazgos:
            filas.append(
                [
                    str(hallazgo.get("clausula_id", "")),
                    str(hallazgo.get("tipo_riesgo", "")),
                    str(hallazgo.get("severidad", "")),
                    str(hallazgo.get("articulo_violado", "")),
                ]
            )
        tabla = Table(filas, hAlign="LEFT")
        tabla.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        elementos.append(tabla)
    else:
        elementos.append(Paragraph("No se detectaron hallazgos de riesgo relevantes.", estilos["Normal"]))

    doc.build(elementos)
    logger.info(f"Reporte PDF generado para contrato {resultado.get('contrato_id')}")
    return buffer.getvalue()

def generar_reporte_xlsx(resultado: dict) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    encabezado_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    encabezado_font = Font(color="FFFFFF", bold=True)

    wb = Workbook()

    resumen = wb.active
    resumen.title = "Resumen"
    resumen.append(["Campo", "Valor"])
    for celda in resumen[1]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font
    filas_resumen = [
        ("Contrato ID", resultado.get("contrato_id")),
        ("Nombre", resultado.get("nombre")),
        ("Categoria", resultado.get("categoria")),
        ("Score", resultado.get("score")),
        ("Nivel de riesgo", resultado.get("nivel_riesgo")),
        ("Resumen ejecutivo", resultado.get("resumen_ejecutivo")),
    ]
    for fila in filas_resumen:
        resumen.append(fila)
    resumen.column_dimensions["A"].width = 20
    resumen.column_dimensions["B"].width = 60

    hallazgos_sheet = wb.create_sheet("Hallazgos")
    encabezados = ["Clausula ID", "Tipo de riesgo", "Severidad", "Articulo violado", "Descripcion"]
    hallazgos_sheet.append(encabezados)
    for celda in hallazgos_sheet[1]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font

    hallazgos = resultado.get("hallazgos", [])
    for hallazgo in hallazgos:
        hallazgos_sheet.append(
            [
                hallazgo.get("clausula_id"),
                hallazgo.get("tipo_riesgo"),
                hallazgo.get("severidad"),
                hallazgo.get("articulo_violado"),
                hallazgo.get("descripcion_riesgo"),
            ]
        )
    for columna, ancho in zip("ABCDE", [12, 22, 12, 18, 50]):
        hallazgos_sheet.column_dimensions[columna].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    logger.info(f"Reporte XLSX generado para contrato {resultado.get('contrato_id')}")
    return buffer.getvalue()


def generar_reporte_comparativo_csv(resultados: list[dict]) -> str:
    df = pd.DataFrame(
        [
            {
                "contrato_id": r.get("contrato_id"),
                "nombre": r.get("nombre"),
                "categoria": r.get("categoria"),
                "score": r.get("score"),
                "nivel_riesgo": r.get("nivel_riesgo"),
                "total_clausulas": r.get("total_clausulas"),
                "clausulas_riesgosas": r.get("clausulas_riesgosas"),
                "motor_riesgo": r.get("motor_riesgo"),
                "fuente_tokenizacion": r.get("fuente_tokenizacion"),
            }
            for r in resultados
        ]
    )
    return df.to_csv(index=False)
